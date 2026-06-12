# src/ui/output_panel.py
# Right panel: OCR raw text, rendered markdown, and extracted fields tabs.

import os
import re
import html
import json
import markdown
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QTabWidget,
                               QPushButton, QLabel, QTextEdit, QMenu,
                               QScrollArea, QFrame, QGridLayout, QApplication,
                               QSizePolicy, QTableWidget, QTableWidgetItem,
                               QProgressBar, QHeaderView, QFileDialog,
                               QPlainTextEdit)
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtWebEngineCore import QWebEnginePage
from PySide6.QtCore import Qt, Slot, QUrl, Signal
from PySide6.QtGui import QTextCursor, QFont


# Dark-themed CSS for the WebEngine rendered output
BROWSER_STYLE = """
<style>
* { box-sizing: border-box; }
body {
    background-color: #0A0A0A;
    color: #F1F5F9;
    font-family: "Be Vietnam Pro", "Noto Sans", "Segoe UI", sans-serif;
    font-size: 11pt;
    line-height: 1.7;
    margin: 16px 20px;
}
a { color: #3B82F6; text-decoration: none; }
a:hover { text-decoration: underline; }
code {
    background-color: #121212;
    color: #93C5FD;
    padding: 2px 6px;
    border-radius: 3px;
    font-family: "Consolas", "Courier New", monospace;
    font-size: 9.5pt;
}
pre {
    background-color: #121212;
    padding: 12px 16px;
    border-radius: 6px;
    overflow-x: auto;
    border: 1px solid #1E293B;
}
pre code { background-color: transparent; padding: 0; color: #F1F5F9; }
h1, h2, h3, h4, h5, h6 { color: #F1F5F9; margin-top: 1.4em; margin-bottom: 0.4em; font-weight: 700; }
h1 { border-bottom: 1px solid #1E293B; padding-bottom: 0.3em; font-size: 1.5em; }
h2 { font-size: 1.25em; }
blockquote {
    border-left: 3px solid #3B82F6;
    margin-left: 0;
    padding-left: 16px;
    color: #94A3B8;
    font-style: italic;
}
ul, ol { padding-left: 24px; }
li { margin-bottom: 0.3em; }
table { border-collapse: collapse; width: 100%; margin-bottom: 1em; }
th, td { border: 1px solid #1E293B; padding: 8px 12px; text-align: left; }
th { background-color: #121212; font-weight: 700; color: #3B82F6; }
tr:nth-child(even) { background-color: #0D0D0D; }
tr:hover { background-color: #1A1A1A; }
hr { border: none; border-top: 1px solid #1E293B; margin: 1.5em 0; }
</style>
"""


def balance_latex_delimiters(latex):
    """Fix unbalanced \\left and \\right delimiters in LaTeX."""
    commands = []
    for m in re.finditer(r'(\\left|\\right)', latex):
        commands.append((m.start(), m.end(), m.group(0)))

    to_remove = set()
    stack = 0

    for start, end, cmd in commands:
        if cmd == r'\left':
            stack += 1
        else:
            if stack > 0:
                stack -= 1
            else:
                to_remove.add(start)

    fixed_latex = ""
    last_idx = 0
    for start, end, cmd in commands:
        if start in to_remove:
            fixed_latex += latex[last_idx:start]
            last_idx = end

    fixed_latex += latex[last_idx:]

    if stack > 0:
        fixed_latex += (r" \right." * stack)

    return fixed_latex


# ==================== Tab 2: Fancy Output ====================
class FancyOutput(QWebEngineView):
    """WebView that renders markdown with MathJax LaTeX support."""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.page().setBackgroundColor(Qt.black)

    def set_markdown(self, md_content):
        if not md_content:
            self.setHtml("")
            return

        try:
            math_blocks = []

            def replace_math(match):
                block = match.group(0)
                block = balance_latex_delimiters(block)
                math_blocks.append(block)
                return f"MATHJAXBLOCKPLACEHOLDER{len(math_blocks)-1}END"

            pattern = re.compile(r'(\\\[.*?\\\])|(\\\(.*?\\\))', re.DOTALL)
            processed_md = pattern.sub(replace_math, md_content)

            html_content = markdown.markdown(processed_md, extensions=['tables', 'fenced_code', 'nl2br'])

            for i, block in enumerate(math_blocks):
                safe_block = html.escape(block)
                html_content = html_content.replace(f"MATHJAXBLOCKPLACEHOLDER{i}END", safe_block)

            base_path = os.path.dirname(os.path.abspath(__file__))
            project_path = os.path.dirname(base_path)
            node_path = os.path.join(project_path, "res", "node")

            mathjax_path = os.path.join(node_path, "mathjax", "tex-mml-svg.js")
            mathjax_url = QUrl.fromLocalFile(mathjax_path).toString()
            mathjax_dir_url = QUrl.fromLocalFile(os.path.dirname(mathjax_path)).toString()
            node_path_url = QUrl.fromLocalFile(node_path).toString()

            newcm_path = os.path.join(node_path, "@mathjax", "mathjax-newcm-font")
            newcm_url = QUrl.fromLocalFile(newcm_path).toString()

            mathjax_script = f"""
            <script>
            MathJax = {{
              loader: {{
                paths: {{
                  mathjax: '{mathjax_dir_url}',
                  npm: '{node_path_url}',
                  'mathjax-newcm-font': '{newcm_url}'
                }},
                pathFilters: [
                  [(data) => {{
                    var cdn = 'https://cdn.jsdelivr.net/npm/';
                    if (data.name.indexOf(cdn) === 0) {{
                        data.name = data.name.replace(cdn, '{node_path_url}/');
                    }}
                    return true;
                  }}, 25]
                ]
              }},
              tex: {{
                inlineMath: [['\\\\(', '\\\\)']],
                displayMath: [['\\\\[', '\\\\]']],
                processEscapes: true
              }},
              svg: {{ fontCache: 'global' }},
              options: {{
                ignoreHtmlClass: 'tex2jax_ignore',
                processHtmlClass: 'tex2jax_process',
                menuOptions: {{ settings: {{ assistiveMml: true, enrich: false }} }}
              }}
            }};
            </script>
            <script id="MathJax-script" async src="{mathjax_url}"></script>
            """

            full_html = f"<html><head>{BROWSER_STYLE}{mathjax_script}</head><body>{html_content}</body></html>"
            self.setHtml(full_html, QUrl.fromLocalFile(project_path))
        except Exception as e:
            print(f"Markdown render error: {e}")

    def copy_content(self):
        self.setFocus()
        page = self.page()
        page.triggerAction(QWebEnginePage.WebAction.SelectAll)
        page.triggerAction(QWebEnginePage.WebAction.Copy)

    def contextMenuEvent(self, event):
        menu = QMenu(self)
        menu.addAction(self.page().action(QWebEnginePage.WebAction.Copy))
        menu.addAction(self.page().action(QWebEnginePage.WebAction.SelectAll))
        menu.exec(event.globalPos())


# ==================== Tab 3: Extracted Fields ====================
class ExtractedFieldsPanel(QWidget):
    """
    Displays extracted document metadata as a clean form-like view.
    Parses JSON blocks from OCR output and shows each field with
    a label/value layout. Loại bản gets a colour-coded badge.
    """

    # Field definitions: (json_key, display_label)
    FIELD_DEFS = [
        ("tac_gia",        "Tác giả"),
        ("the_loai",       "Thể loại"),
        ("ngay_thang_nam", "Ngày"),
        ("trich_yeu",      "Trích yếu"),
        ("do_mat",    "Độ mật"),
        ("nguoi_ky",  "Người ký"),
        ("loai_ban",  "Loại bản"),
    ]

    # Badge mappings: value substring -> object_name
    BADGE_MAP = {
        "bản gốc":  "badge_original",
        "bản chính": "badge_main",
        "bản photo": "badge_photo",
        "photo":     "badge_photo",
    }

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("extracted_panel")

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        # Scrollable content area
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)

        content_widget = QWidget()
        content_widget.setObjectName("extracted_panel")
        self._grid = QGridLayout(content_widget)
        self._grid.setContentsMargins(16, 16, 16, 16)
        self._grid.setHorizontalSpacing(16)
        self._grid.setVerticalSpacing(10)
        self._grid.setColumnStretch(1, 1)

        self._value_labels: dict[str, QLabel] = {}
        self._build_fields()

        scroll.setWidget(content_widget)
        outer.addWidget(scroll)

        # Bottom bar: Copy JSON button
        bottom_bar = QHBoxLayout()
        bottom_bar.setContentsMargins(16, 8, 16, 8)
        bottom_bar.addStretch()

        self.btn_copy_json = QPushButton("Sao chép JSON")
        self.btn_copy_json.setObjectName("btn_copy_json")
        self.btn_copy_json.setFixedHeight(32)
        self.btn_copy_json.clicked.connect(self._copy_json)
        bottom_bar.addWidget(self.btn_copy_json)

        outer.addLayout(bottom_bar)

        # Store last parsed data for JSON copy
        self._current_data: dict = {}

    def _build_fields(self):
        for row, (key, label_text) in enumerate(self.FIELD_DEFS):
            lbl = QLabel(label_text + ":")
            lbl.setObjectName("field_label")
            lbl.setAlignment(Qt.AlignTop | Qt.AlignRight)
            lbl.setSizePolicy(QSizePolicy.Minimum, QSizePolicy.Preferred)
            lbl.setMinimumWidth(90)

            val = QLabel("—")
            val.setObjectName("field_value")
            val.setProperty("empty", "true")
            val.setWordWrap(True)
            val.setTextInteractionFlags(Qt.TextSelectableByMouse)
            val.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)

            self._grid.addWidget(lbl, row, 0)
            self._grid.addWidget(val, row, 1)
            self._value_labels[key] = val

        # Push rows to top
        self._grid.setRowStretch(len(self.FIELD_DEFS), 1)

    def _make_badge(self, text: str) -> QLabel:
        badge = QLabel(text)
        badge.setTextInteractionFlags(Qt.TextSelectableByMouse)
        key = text.strip().lower()
        obj_name = "badge_unknown"
        for fragment, name in self.BADGE_MAP.items():
            if fragment in key:
                obj_name = name
                break
        badge.setObjectName(obj_name)
        badge.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        return badge

    def populate(self, data: dict):
        """Populate all field labels from a dict of extracted values."""
        self._current_data = data

        for key, label_text in self.FIELD_DEFS:
            val_widget = self._value_labels[key]
            raw_value = data.get(key, "")
            text = str(raw_value).strip() if raw_value else ""

            if key == "loai_ban" and text:
                # Replace the plain label with a badge label
                row = list(k for k, _ in self.FIELD_DEFS).index(key)
                old = self._grid.itemAtPosition(row, 1)
                if old and old.widget():
                    old.widget().hide()
                    self._grid.removeWidget(old.widget())

                badge = self._make_badge(text)
                self._grid.addWidget(badge, row, 1)
                self._value_labels[key] = badge
            else:
                if text:
                    val_widget.setText(text)
                    val_widget.setProperty("empty", "false")
                else:
                    val_widget.setText("—")
                    val_widget.setProperty("empty", "true")

                # Force style refresh
                val_widget.style().unpolish(val_widget)
                val_widget.style().polish(val_widget)

    def clear(self):
        """Reset all fields to empty state."""
        self._current_data = {}
        for key, _ in self.FIELD_DEFS:
            row = list(k for k, _ in self.FIELD_DEFS).index(key)
            old = self._grid.itemAtPosition(row, 1)
            if old and old.widget():
                old.widget().hide()
                self._grid.removeWidget(old.widget())

            val = QLabel("—")
            val.setObjectName("field_value")
            val.setProperty("empty", "true")
            val.setWordWrap(True)
            val.setTextInteractionFlags(Qt.TextSelectableByMouse)
            val.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
            val.show()
            self._grid.addWidget(val, row, 1)
            self._value_labels[key] = val

    def _copy_json(self):
        if not self._current_data:
            return
        text = json.dumps(self._current_data, ensure_ascii=False, indent=2)
        QApplication.clipboard().setText(text)

    def try_parse_and_populate(self, raw_text: str) -> bool:
        """
        Attempt to extract a JSON block from raw OCR text and populate fields.
        Returns True if parsing succeeded.
        """
        # Try to find a JSON object in the text
        match = re.search(r'\{[^{}]*\}', raw_text, re.DOTALL)
        if not match:
            # Try fenced code block containing JSON
            match = re.search(r'```(?:json)?\s*(\{.*?\})\s*```', raw_text, re.DOTALL)
            if match:
                json_str = match.group(1)
            else:
                return False
        else:
            json_str = match.group(0)

        try:
            data = json.loads(json_str)
            if isinstance(data, dict):
                self.populate(data)
                return True
        except json.JSONDecodeError:
            pass

        return False


# ==================== Tab 4: Batch Results ====================
BATCH_COLUMNS = ["STT", "Tên file", "Tác giả", "Thể loại", "Ngày", "Người ký", "Độ mật", "Trạng thái"]


class BatchResultsPanel(QWidget):
    """Tab showing real-time batch scan results with CSV/Excel export."""

    def __init__(self, parent=None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(6)

        self._progress = QProgressBar()
        self._progress.setTextVisible(True)
        self._progress.setFormat("Đang scan %v/%m file...")
        self._progress.hide()
        layout.addWidget(self._progress)

        self._status_label = QLabel("—")
        self._status_label.setStyleSheet("font-weight: bold; padding: 2px 0;")
        layout.addWidget(self._status_label)

        self._log_area = QPlainTextEdit()
        self._log_area.setReadOnly(True)
        self._log_area.setFixedHeight(80)
        self._log_area.setFont(QFont("Consolas", 9))
        self._log_area.setPlaceholderText("Log scan sẽ hiển thị ở đây...")
        layout.addWidget(self._log_area)

        self._table = QTableWidget(0, len(BATCH_COLUMNS))
        self._table.setHorizontalHeaderLabels(BATCH_COLUMNS)
        self._table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self._table.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        self._table.horizontalHeader().setSectionResizeMode(4, QHeaderView.Stretch)
        self._table.setEditTriggers(QTableWidget.NoEditTriggers)
        self._table.setSelectionBehavior(QTableWidget.SelectRows)
        self._table.setAlternatingRowColors(True)
        layout.addWidget(self._table)

        btn_bar = QHBoxLayout()
        btn_bar.addStretch()

        self.btn_export_csv = QPushButton("Xuất CSV")
        self.btn_export_csv.setFixedHeight(32)
        self.btn_export_csv.setEnabled(False)
        self.btn_export_csv.clicked.connect(self._export_csv)
        btn_bar.addWidget(self.btn_export_csv)

        self.btn_export_excel = QPushButton("Xuất Excel")
        self.btn_export_excel.setFixedHeight(32)
        self.btn_export_excel.setEnabled(False)
        self.btn_export_excel.clicked.connect(self._export_excel)
        btn_bar.addWidget(self.btn_export_excel)

        layout.addLayout(btn_bar)

        self._rows: list = []

    def start_scan(self, total: int):
        """Call before BatchScanWorker starts."""
        self._table.setRowCount(0)
        self._rows.clear()
        self.btn_export_csv.setEnabled(False)
        self.btn_export_excel.setEnabled(False)
        self._progress.setMaximum(total)
        self._progress.setValue(0)
        self._progress.show()
        self._log_area.clear()
        self._status_label.setText("—")

    def append_row(self, result: dict):
        """Append one scan result row. Called from row_ready signal handler."""
        self._rows.append(result)
        row_idx = self._table.rowCount()
        self._table.insertRow(row_idx)

        fields = result.get("fields", {})
        status = result.get("status", "error")

        values = [
            str(row_idx + 1),
            result.get("filename", ""),
            fields.get("tac_gia", "—") or "—",
            fields.get("the_loai", "—") or "—",
            fields.get("ngay_thang_nam", "—") or "—",
            fields.get("nguoi_ky", "—") or "—",
            fields.get("do_mat", "—") or "—",
            "✓ OK" if status == "ok" else f"✗ {result.get('error_msg', 'Lỗi')}",
        ]

        for col, val in enumerate(values):
            item = QTableWidgetItem(val)
            item.setTextAlignment(Qt.AlignVCenter | Qt.AlignLeft)
            self._table.setItem(row_idx, col, item)

        self._progress.setValue(row_idx + 1)
        self._table.scrollToBottom()

    def finish_scan(self):
        """Call when BatchScanWorker emits finished."""
        self._progress.hide()
        if self._rows:
            self.btn_export_csv.setEnabled(True)
            self.btn_export_excel.setEnabled(True)

    def append_log(self, text: str):
        self._log_area.moveCursor(QTextCursor.End)
        self._log_area.insertPlainText(text + "\n")
        self._log_area.moveCursor(QTextCursor.End)

    def update_status(self, text: str):
        self._status_label.setText(text)

    def _export_csv(self):
        import csv
        path, _ = QFileDialog.getSaveFileName(self, "Xuất CSV", "batch_results.csv",
                                               "CSV Files (*.csv)")
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(BATCH_COLUMNS)
            for i, r in enumerate(self._rows):
                fields = r.get("fields", {})
                writer.writerow([
                    i + 1,
                    r.get("filename", ""),
                    fields.get("tac_gia", ""),
                    fields.get("the_loai", ""),
                    fields.get("ngay_thang_nam", ""),
                    fields.get("nguoi_ky", ""),
                    fields.get("do_mat", ""),
                    "OK" if r.get("status") == "ok" else r.get("error_msg", "Lỗi"),
                ])

    def _export_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            from PySide6.QtWidgets import QMessageBox
            QMessageBox.critical(self, "Lỗi", "Chưa cài openpyxl. Chạy: pip install openpyxl")
            return

        path, _ = QFileDialog.getSaveFileName(self, "Xuất Excel", "batch_results.xlsx",
                                               "Excel Files (*.xlsx)")
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kết quả Batch"

        header_fill = PatternFill("solid", fgColor="1E293B")
        header_font = Font(bold=True, color="93C5FD")
        ws.append(BATCH_COLUMNS)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for i, r in enumerate(self._rows):
            fields = r.get("fields", {})
            ws.append([
                i + 1,
                r.get("filename", ""),
                fields.get("tac_gia", ""),
                fields.get("the_loai", ""),
                fields.get("ngay_thang_nam", ""),
                fields.get("nguoi_ky", ""),
                fields.get("do_mat", ""),
                "OK" if r.get("status") == "ok" else r.get("error_msg", "Lỗi"),
            ])

        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 25
        wb.save(path)


# ==================== Main Widget ====================
class OutputPanel(QWidget):
    # Emitted when user clicks "Trích xuất" — carries current OCR text
    extract_requested = Signal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("output_panel_widget")
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        self.t = {}

        # === Tabs ===
        self.tabs = QTabWidget()
        self.tabs.setDocumentMode(False)

        # --- Tab 1: Raw OCR text (streaming) ---
        self.tab_raw = QWidget()
        raw_layout = QVBoxLayout(self.tab_raw)
        raw_layout.setContentsMargins(0, 0, 0, 0)
        raw_layout.setSpacing(0)

        self.text_output = QTextEdit()
        self.text_output.setReadOnly(True)
        raw_layout.addWidget(self.text_output)

        # --- Tab 2: Rendered markdown (fancy) ---
        self.web_view = FancyOutput()

        # --- Tab 3: Extracted fields ---
        self.extracted_panel = ExtractedFieldsPanel()

        self.tabs.addTab(self.tab_raw, "")       # index 0
        self.tabs.addTab(self.web_view, "")      # index 1
        self.tabs.addTab(self.extracted_panel, "") # index 2

        # --- Tab 4: Batch scan results ---
        self.batch_results_panel = BatchResultsPanel()
        self.tabs.addTab(self.batch_results_panel, "Kết quả Batch")  # index 3
        self.tabs.setTabEnabled(3, False)

        self.tabs.setTabEnabled(1, False)
        self.tabs.setTabEnabled(2, False)
        self.tabs.currentChanged.connect(self._update_copy_button_text)

        # === Bottom bar ===
        bottom_bar = QHBoxLayout()
        bottom_bar.setContentsMargins(8, 6, 8, 6)
        bottom_bar.setSpacing(8)

        self.lbl_proofread = QLabel()
        self.lbl_proofread.setObjectName("lbl_proofread")
        self.lbl_proofread.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)

        self.btn_extract = QPushButton("Trích xuất fields")
        self.btn_extract.setObjectName("btn_extract")
        self.btn_extract.setFixedHeight(32)
        self.btn_extract.setEnabled(False)
        self.btn_extract.clicked.connect(self._on_extract_clicked)

        self.btn_copy = QPushButton()
        self.btn_copy.setObjectName("btn_copy")
        self.btn_copy.setFixedHeight(32)
        self.btn_copy.clicked.connect(self.copy_output)

        bottom_bar.addWidget(self.lbl_proofread, stretch=1)
        bottom_bar.addWidget(self.btn_extract)
        bottom_bar.addWidget(self.btn_copy)

        main_layout.addWidget(self.tabs)
        main_layout.addLayout(bottom_bar)

    # ==================== Language ====================
    def update_language(self, t):
        self.t = t
        self.lbl_proofread.setText(t.get("lbl_proofread", ""))
        self.tabs.setTabText(0, t.get("tab_raw", "Văn bản OCR"))
        self.tabs.setTabText(1, t.get("tab_fancy", "Kết quả đẹp"))
        self.tabs.setTabText(2, t.get("tab_extracted", "Trích xuất"))
        self.tabs.setTabText(3, t.get("tab_batch", "Kết quả Batch"))
        self._update_copy_button_text()

    def _update_copy_button_text(self):
        if not self.t or "btn_copy" not in self.t:
            return
        current_tab_text = self.tabs.tabText(self.tabs.currentIndex())
        self.btn_copy.setText(self.t["btn_copy"].format(current_tab_text))

    def _on_extract_clicked(self):
        text = self.text_output.toPlainText().strip()
        if text:
            self.extract_requested.emit(text)

    # ==================== Tab 1: Raw Output ====================
    @Slot(str)
    def append_text(self, text):
        """Append streaming text. Always switches to raw tab while streaming."""
        if self.tabs.currentIndex() != 0:
            self.tabs.setCurrentIndex(0)

        if self.tabs.isTabEnabled(1):
            self.tabs.setTabEnabled(1, False)
        if self.tabs.isTabEnabled(2):
            self.tabs.setTabEnabled(2, False)

        self.text_output.moveCursor(QTextCursor.End)
        self.text_output.insertPlainText(text)
        self.text_output.moveCursor(QTextCursor.End)

    # ==================== Tab 2: Fancy Output ====================
    def render_fancy_output(self):
        """Render OCR text as markdown and enable the extract button."""
        raw_md = self.text_output.toPlainText()
        if not raw_md.strip():
            return

        self.web_view.set_markdown(raw_md)
        self.tabs.setTabEnabled(1, True)
        self.btn_extract.setEnabled(True)

        # Switch to fancy tab by default
        self.tabs.setCurrentIndex(1)

    # ==================== Utility ====================
    def set_corrected_text(self, corrected_text):
        # Replace raw output content with corrected text (preserves appended headers).
        self.text_output.clear()
        self.text_output.setPlainText(corrected_text)

    def clear(self):
        self.text_output.clear()
        self.web_view.set_markdown("")
        self.extracted_panel.clear()
        self.tabs.setTabEnabled(1, False)
        self.tabs.setTabEnabled(2, False)
        self.tabs.setCurrentIndex(0)
        self.btn_extract.setEnabled(False)

    def copy_output(self):
        idx = self.tabs.currentIndex()
        if idx == 0:
            self.text_output.selectAll()
            self.text_output.copy()
        elif idx == 1:
            self.web_view.copy_content()
        elif idx == 2:
            self.extracted_panel._copy_json()
